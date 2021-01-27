from flask import make_response
from openpyxl import Workbook
from openpyxl.writer.excel import save_virtual_workbook

from app.models.therapy_record import TreRec
from app.utils.date import get_birth_date_by_id_card, get_age_by_birth


class Export:
    lab_inspectation_table = ['BloodRoutine', 'BloodBio', 'Thyroid', 'Coagulation', 'MyocardialEnzyme',
                              'Cytokines', 'LymSubsets', 'UrineRoutine', 'TumorMarker']
    other_inspect_table = ['Lung', 'OtherExams', 'ImageExams']
    header_map = {'patNumber': '编号', 'researchCenter':'研究中心', 'idNumber': '身份证号', 'patientName': '姓名', 'hospitalNumber': '住院号',
                  'gender': '性别', 'birthday': '出生日期', 'age': '年龄', 'phoneNumber1': '电话号码1', 'phoneNumber2': '电话号码2',
                  'basDisHis': '基础疾病史', 'infDisHis': '传染疾病史', 'tumor': '是否有肿瘤史', 'tumHis': '肿瘤史','tumorFam': '是否有肿瘤家族史',
                  'tumFamHis': '肿瘤家族史', 'smoke': '是否吸烟', 'smokingHis': '吸烟史', 'drink': '是否饮酒', 'drinkingHis': '饮酒史',
                  'hormone': '是否长期使用激素', 'hormoneUseHis': '激素使用史', 'drug': '是否长期使用药物', 'drugUseHis': '药物使用史',
                  'PSScore': 'PS评分', 'cliniManifest': '临床表现', 'videography': '影像学', 'part': '部位', 'bioMet': '活检方式', 'pleInv': '是否胸膜侵犯',
                  'speSite': '标本部位', 'firVisDate': '初诊日期', 'patReDate': '病理报告日期', 'patNum': '病理号', 'patDia': '病理诊断', 'mitIma': '核分裂像',
                  'comCar': '复合性癌', 'necArea': '坏死面积', 'massSize': '肿块大小', 'Ki67': 'Ki67', 'traSite': '转移部位', 'TSize': 'TSize', 'stage': '分期情况',
                  'cliStage': '临床分期', 'patStage': '病理分期', 'cRemark': 'C备注', 'pRemark': 'P备注',
                  'samplingTime': '采样时间', 'detectTime': '检测时间',
                  'ECGDetTime': '12导联心电图检测时间', 'ECGDesc': '12导联心电图结果描述', 'ECGPath': '12导联心电图报告路径',
                  'UCGDetTime': '超声心动图检测时间', 'UCGDesc': '超声心动图结果描述', 'UCGPath': '超声心动图报告路径',
                  'detectTime': '检测时间', 'examArea': '检查部位', 'exmaMethod': '检查方法', 'tumorLD': '肿瘤长径', 'tumorSD': '肿瘤短径',
                  'tumorDesc': '肿瘤描述', 'path': '文件路径',
                  'other': '其他', 'filePath': '文件路径', 'PDL1': 'PD-L1表达水平', 'PDL1KT': 'PDL1抗体', 'PD1':'PD1表达水平', 'PD1KT':'PD1抗体',
                  'trement': '几线治疗','beEffEvaDate':'最佳疗效评估日期','beEffEva':'最佳疗效评估','proDate':'进展日期',
                  'proDes':'进展描述','PFS_DFS':'PFS/DFS',
                  'isTre': '是否加入临床治疗', 'clinTri': '临床实验名称', 'treSolu': '治疗方案', 'begDate': '开始日期',
                  'endDate': '结束日期', 'isRepBio': '是否重复活检', 'bioMet': '活检方式', 'matPart': '取材部位',
                  'specNum': '标本库流水号', 'patDia': '病理诊断',
                  'surSco': '手术范围', 'lymDis': '淋巴清扫范围', 'cleGro': '清扫组数', 'surDate': '手术日期', 'posAdjChem': '术后辅助化疗',
                  'isRepBio': '是否重复活检', 'bioMet': '活检方式', 'matPart': '取材部位', 'specNum': '标本库流水号', 'patDia': '病理诊断',
                  'begDate': '开始日期', 'endDate': '结束日期', 'radSite': '放疗部位', 'radDose': '放疗剂量', 'splTim': '分割次数',
                  'isRepBio': '是否重复活检', 'bioMet': '活检方式', 'matPart': '取材部位', 'specNum': '标本库流水号', 'patDia': '病理诊断'

                  }
    treRec_header_map = {'trement': '几线治疗','beEffEvaDate':'最佳疗效评估日期','beEffEva':'最佳疗效评估','proDate':'进展日期',
                  'proDes':'进展描述','PFS_DFS':'PFS/DFS'}
    oneToFive_header_map = {'isTre':'是否加入临床治疗', 'clinTri':'临床实验名称', 'treSolu':'治疗方案', 'begDate':'开始日期',
                     'endDate':'结束日期', 'isRepBio':'是否重复活检', 'bioMet':'活检方式', 'matPart':'取材部位',
                     'specNum':'标本库流水号', 'patDia':'病理诊断'}
    surgery_header_map = {'surSco':'手术范围', 'lymDis':'淋巴清扫范围', 'cleGro':'清扫组数', 'surDate':'手术日期', 'posAdjChem':'术后辅助化疗',
                          'isRepBio':'是否重复活检', 'bioMet':'活检方式', 'matPart':'取材部位', 'specNum':'标本库流水号','patDia':'病理诊断'}
    radio_header_map = {'begDate':'开始日期', 'endDate':'结束日期', 'radSite':'放疗部位', 'radDose':'放疗剂量', 'splTim':'分割次数',
                        'isRepBio':'是否重复活检', 'bioMet':'活检方式', 'matPart':'取材部位', 'specNum':'标本库流水号','patDia':'病理诊断'}

    patDia_map = {
        "0-0": "上皮型肿瘤",
        "0-0-0": "上皮型肿瘤-腺癌",
        "0-0-0-0": "上皮型肿瘤-腺癌-贴壁型腺癌",
        "0-0-0-1": "上皮型肿瘤-腺癌-腺泡型腺癌",
        "0-0-0-2": "上皮型肿瘤-腺癌-乳头型腺癌",
        "0-0-0-3": "上皮型肿瘤-腺癌-微乳头型腺癌",
        "0-0-0-4": "上皮型肿瘤-腺癌-实体型腺癌",
        "0-0-0-5": "上皮型肿瘤-腺癌-浸润性粘液腺癌",
        "0-0-0-5-0": "上皮型肿瘤-腺癌-浸润性粘液腺癌-浸润和非浸润性混合型粘液性腺癌",
        "0-0-0-6": "上皮型肿瘤-腺癌-胶质性腺癌",
        "0-0-0-7": "上皮型肿瘤-腺癌-胎儿型腺癌",
        "0-0-0-8": "上皮型肿瘤-腺癌-肠型腺癌",
        "0-0-0-9": "上皮型肿瘤-腺癌-微小浸润性腺癌",
        "0-0-0-9-0": "上皮型肿瘤-腺癌-微小浸润性腺癌-非粘液性腺癌",
        "0-0-0-9-1": "上皮型肿瘤-腺癌-微小浸润性腺癌-粘液癌",
        "0-0-0-10": "上皮型肿瘤-腺癌-侵袭前病变",
        "0-0-0-10-0": "上皮型肿瘤-腺癌-侵袭前病变-非典型腺瘤样增生",
        "0-0-0-10-1": "上皮型肿瘤-腺癌-侵袭前病变-原位腺癌",
        "0-0-0-10-1-0": "上皮型肿瘤-腺癌-侵袭前病变-原位腺癌-非粘液性原位腺癌",
        "0-0-0-10-1-1": "上皮型肿瘤-腺癌-侵袭前病变-原位腺癌-粘液性原位腺癌",
        "0-0-1": "上皮型肿瘤-鳞癌",
        "0-0-1-0": "上皮型肿瘤-鳞癌-角化型鳞癌",
        "0-0-1-1": "上皮型肿瘤-鳞癌-非角化型鳞癌",
        "0-0-1-2": "上皮型肿瘤-鳞癌-基底鳞状细胞癌",
        "0-0-1-3": "上皮型肿瘤-鳞癌-侵袭前病变",
        "0-0-1-3-0": "上皮型肿瘤-鳞癌-侵袭前病变-鳞状细胞原位癌",
        "0-1": "神经内分泌肿瘤",
        "0-1-0": "神经内分泌肿瘤-小细胞癌",
        "0-1-0-0": "神经内分泌肿瘤-小细胞癌-结合小细胞癌",
        "0-1-1": "神经内分泌肿瘤-大细胞神经内分泌癌",
        "0-1-1-0": "神经内分泌肿瘤-大细胞神经内分泌癌-结合大细胞神经内分泌癌",
        "0-1-2": "神经内分泌肿瘤-类癌肿瘤",
        "0-1-2-0": "神经内分泌肿瘤-类癌肿瘤-典型类癌肿瘤",
        "0-1-2-1": "神经内分泌肿瘤-类癌肿瘤-非典型类癌肿瘤",
        "0-1-3": "神经内分泌肿瘤-侵袭前的病变",
        "0-1-3-0": "神经内分泌肿瘤-侵袭前的病变-弥漫性特发性肺神经内分泌细胞增生",
        "0-1-4": "神经内分泌肿瘤-大细胞癌",
        "0-1-5": "神经内分泌肿瘤-腺鳞癌",
        "0-1-6": "神经内分泌肿瘤-癌肉瘤样癌",
        "0-1-6-0": "神经内分泌肿瘤-癌肉瘤样癌-多形性癌",
        "0-1-6-1": "神经内分泌肿瘤-癌肉瘤样癌-梭形细胞癌",
        "0-1-6-2": "神经内分泌肿瘤-癌肉瘤样癌-巨细胞癌",
        "0-1-6-3": "神经内分泌肿瘤-癌肉瘤样癌-癌肉瘤",
        "0-1-6-4": "神经内分泌肿瘤-癌肉瘤样癌-肺胚细胞瘤",
        "0-1-7": "神经内分泌肿瘤-其他未分类癌",
        "0-1-7-0": "神经内分泌肿瘤-其他未分类癌-淋巴上皮样癌",
        "0-1-7-1": "神经内分泌肿瘤-其他未分类癌-NUT肿瘤",
        "0-1-8": "神经内分泌肿瘤-唾液型肿瘤",
        "0-1-8-0": "神经内分泌肿瘤-唾液型肿瘤-粘液表皮样癌肿瘤",
        "0-1-8-1": "神经内分泌肿瘤-唾液型肿瘤-腺样囊性癌",
        "0-1-8-2": "神经内分泌肿瘤-唾液型肿瘤-上皮-肌上皮癌",
        "0-1-8-3": "神经内分泌肿瘤-唾液型肿瘤-多形性腺瘤",
        "0-1-9": "神经内分泌肿瘤-乳头状瘤",
        "0-1-9-0": "神经内分泌肿瘤-乳头状瘤-鳞状细胞乳头状癌",
        "0-1-9-0-0": "神经内分泌肿瘤-乳头状瘤-鳞状细胞乳头状癌-外生型",
        "0-1-9-0-1": "神经内分泌肿瘤-乳头状瘤-鳞状细胞乳头状癌-逆向生长",
        "0-1-9-1": "神经内分泌肿瘤-乳头状瘤-腺型状瘤",
        "0-1-9-2": "神经内分泌肿瘤-乳头状瘤-腺鳞混合型乳头状瘤",
        "0-1-10": "神经内分泌肿瘤-腺瘤",
        "0-1-10-0": "神经内分泌肿瘤-腺瘤-良性硬化性肺细胞瘤",
        "0-1-10-1": "神经内分泌肿瘤-腺瘤-泡腺腺瘤",
        "0-1-10-2": "神经内分泌肿瘤-腺瘤-乳头状腺瘤",
        "0-1-10-3": "神经内分泌肿瘤-腺瘤-粘液性囊腺瘤腺瘤",
        "0-1-10-4": "神经内分泌肿瘤-腺瘤-粘液腺腺瘤",
        "0-2": "间叶性肿瘤",
        "0-2-0": "间叶性肿瘤-肺错构瘤",
        "0-2-1": "间叶性肿瘤-软骨瘤",
        "0-2-2": "间叶性肿瘤-PEComatous肿瘤",
        "0-2-2-0": "间叶性肿瘤-PEComatous肿瘤-淋巴管平滑肌瘤病",
        "0-2-2-1": "间叶性肿瘤-PEComatous肿瘤-PEComa-良性",
        "0-2-2-1-0": "间叶性肿瘤-PEComatous肿瘤-PEComa-良性-透明细胞瘤",
        "0-2-2-2": "间叶性肿瘤-PEComatous肿瘤-PEComa-恶性",
        "0-2-3": "间叶性肿瘤-先天性支气管周肌纤维母细胞肿瘤",
        "0-2-4": "间叶性肿瘤-弥漫性肺淋巴管瘤病",
        "0-2-5": "间叶性肿瘤-炎症性肌纤维母细胞瘤",
        "0-2-6": "间叶性肿瘤-上皮样血管内皮瘤",
        "0-2-7": "间叶性肿瘤-胸膜肺母细胞瘤",
        "0-2-8": "间叶性肿瘤-滑膜肉瘤",
        "0-2-9": "间叶性肿瘤-肺动脉内膜肉瘤",
        "0-2-10": "间叶性肿瘤-肺黏液肉瘤伴EWSR1-CREB1易位",
        "0-2-11": "间叶性肿瘤-肌上皮肿瘤",
        "0-2-11-0": "间叶性肿瘤-肌上皮肿瘤-肌上皮瘤",
        "0-2-11-1": "间叶性肿瘤-肌上皮肿瘤-肌上皮癌",
        "0-2-12": "间叶性肿瘤-淋巴细胞组织细胞肿瘤",
        "0-2-13": "间叶性肿瘤-结外边缘区黏膜相关淋巴组织淋巴瘤（MALT淋巴瘤）",
        "0-2-14": "间叶性肿瘤-弥漫性大细胞淋巴瘤",
        "0-2-15": "间叶性肿瘤-淋巴瘤样肉芽肿",
        "0-2-16": "间叶性肿瘤-血管内大B细胞淋巴瘤",
        "0-2-17": "间叶性肿瘤-肺朗格罕细胞组织细胞增生症",
        "0-2-18": "间叶性肿瘤-Erdheim-Chester病",
        "0-3": "异位肿瘤",
        "0-3-0": "异位肿瘤-生殖细胞肿瘤",
        "0-3-0-0": "异位肿瘤-生殖细胞肿瘤-畸胎瘤-成熟",
        "0-3-0-1": "异位肿瘤-生殖细胞肿瘤-畸胎瘤-不成熟",
        "0-3-1": "异位肿瘤-肺内的胸腺瘤",
        "0-3-2": "异位肿瘤-黑色素瘤",
        "0-3-3": "异位肿瘤-脑膜瘤",
        "0-4": "转移性肿瘤",
        "0-5": "其他"
    }  # 病理诊断map
    gender_map = {0: "女", 1: "男", "/": "/"}
    videography_map = {0:'周围型', 1:'中央型', "/": "/"}
    stage_map = {'1':'未住院', '2':'C/P/S均无法分期', '3':'仅C分期', '4':'仅P分期', '5':'C分期和P分期', "/": "/"}
    immunohis_map = {0: '无', 1: '-', 2: '±', 3: '+', 4: '++', 5: '+++', "/": "/"}

    moleDetec_map = {0: '阴性', 1: '阳性', 2: '无', "/": "/"}
    moleDetec_fields = ['ALK', 'BIM', 'BRAF', 'cMET', 'EGFR', 'HER_2', 'KRAS',
                        'PIK3CA', 'ROS1', 'RET', 'UGT1A1', 'NTRK']
    lung_fields = ['FVC', 'FEV1_FVC', 'MEF', 'MEF25', 'MEF50', 'MEF75', 'TLC_sb',
                   'RV', 'RV_TLC', 'VC', 'DLCO_ex', 'DLCO_sb', 'KCO']
    moleDetec_DetMed_map = {1: 'ARMS', 2: 'FISH', 3: 'NGS', "/": "/"}
    moleDetec_MSI_map = {0: 'MSS', 1: 'MSIH', 2: 'MSIL', "/": "/"}
    Mea_map = {-1:'异常', 0: '正常', 1: '异常 1', 2: '异常 2', 3: '异常 3', 4: '异常 4', 5: '异常 5', "/": "/"}
    radosUnit_map = {0:'Gy', 1:'cGy', "/": "/"}
    isTre_map = {0:'否', 1:'是', -1:'不详', "/": "/"}
    isExe_map = {0:'否', 1:'是'}


    def __init__(self, info, tables, columns):
        self.headers = []
        self.info = info
        self.tables = tables
        self.columns = columns

        i = 0
        for x in columns:
            name = tables[i].__name__
            if (name in self.lab_inspectation_table or name == 'Lung'
            or name == 'Immunohis' or name == 'MoleDetec'):
                for y in x:
                    if (y == 'samplingTime' or y == 'other' or y == 'filePath' or y == 'path'
                    or y == 'PDL1' or y == 'PDL1KT' or y == 'PD1' or y == 'PD1KT' or y == 'detectTime'):
                        self.headers.append(self.header_map.get(y))
                    elif y[-3:] == 'Mea':
                        self.headers.append(y[:-3] + '临床意义判断')
                    elif y[-4:] == 'Note':
                        self.headers.append(y[:-4] + '备注')
                    elif y[-3:] == 'Sam':
                        self.headers.append(y[:-3] + '检测样本')
                    elif y[-6:] == 'DetMed':
                        self.headers.append(y[:-6] + '检测方法')
                    elif y[-4:] == 'Desc':
                        self.headers.append(y[:-4] + '检测描述')
                    elif y in self.lung_fields:
                        self.headers.append(y)
                        self.headers.append(y + '最佳值')
                        self.headers.append(y + '最佳值/预期值(%)')
                    elif y in self.header_map.keys():
                        self.headers.append(self.header_map.get(y))
                    else:
                        self.headers.append(y)

            else:
                for y in x:
                    if y == 'cStage':
                        self.headers.append('c分期T')
                        self.headers.append('c分期N')
                        self.headers.append('c分期M')
                    elif y == 'pStage':
                        self.headers.append('p分期T')
                        self.headers.append('p分期N')
                        self.headers.append('p分期M')
                    elif y in self.header_map.keys():
                        self.headers.append(self.header_map.get(y))
                    else:
                        self.headers.append(y)
            i += 1

    def work(self):
        wb = Workbook()
        ws = wb.active
        ws.title = '导出样本数据'
        ws.append(self.headers)

        for dic in self.info:

            pid = dic.get('pid')
            treNum = dic.get('treNum')
            treNums = []
            if treNum > 0:
                for k in range(1, treNum + 1):
                    treNums.append(k)
                for treNum in treNums:
                    row = []
                    i = 0
                    for obj_class_name in self.tables:
                        name = obj_class_name.__name__
                        if name == 'Patient':
                            obj_array = obj_class_name.query.filter(obj_class_name.id == pid,
                                                                    obj_class_name.is_delete == 0).all()
                        elif name == 'PastHis' or name == 'IniDiaPro':
                            obj_array = obj_class_name.query.filter(obj_class_name.pid == pid,
                                                                    obj_class_name.is_delete == 0).all()
                        else:
                            obj_array = obj_class_name.query.filter(obj_class_name.pid == pid,
                                                                    obj_class_name.treNum == treNum,
                                                                    obj_class_name.is_delete == 0).all()

                        if obj_array:
                            obj = obj_array[0]
                            for field in self.columns[i]:
                                if field != 'age' and not(field in self.lung_fields):
                                    x = getattr(obj, field)

                                if name == 'Immunohis':
                                    if field == 'Ki67' or field == 'other' or field == 'filePath':
                                        value = self.filter_none(obj, field)
                                        row.append(value)
                                    else:
                                        value = self.filter_none(obj, field)
                                        value = self.immunohis_map.get(value)
                                        row.append(value)
                                elif name == 'MoleDetec':
                                    if field[-6:] == 'DetMed':
                                        value = self.filter_none(obj, field)
                                        value = self.moleDetec_DetMed_map.get(value)
                                        row.append(value)
                                    elif field == 'MSI':
                                        value = self.filter_none(obj, field)
                                        value = self.moleDetec_MSI_map.get(value)
                                        row.append(value)
                                    elif field in self.moleDetec_fields:
                                        value = self.filter_none(obj, field)
                                        value = self.moleDetec_map.get(value)
                                        row.append(value)
                                    else:
                                        value = self.filter_none(obj, field)
                                        row.append(value)
                                elif name in self.lab_inspectation_table or name == 'Lung':
                                    if field[-3:] == 'Mea':
                                        value = self.filter_none(obj, field)
                                        value = self.Mea_map.get(value)
                                        row.append(value)
                                    elif field in self.lung_fields:
                                        value_exp = self.filter_none(obj, field+'_exp')
                                        value_best = self.filter_none(obj, field+'_best')
                                        value_ratio = self.filter_none(obj, field+'_ratio')
                                        row.extend([value_exp, value_best, value_ratio])
                                    else:
                                        value = self.filter_none(obj, field)
                                        row.append(value)
                                elif field == 'gender':
                                    value = self.filter_none(obj, field)
                                    value = self.gender_map.get(value)
                                    row.append(value)
                                elif field == 'age':
                                    idNumber = getattr(obj, 'idNumber')
                                    age = get_age_by_birth(get_birth_date_by_id_card(idNumber))
                                    value = age if age else '/'
                                    row.append(value)
                                elif field == 'drinkingHis':
                                    value = self.format_drink_history(obj)
                                    row.append(value)
                                elif field == 'smokingHis':
                                    value = self.format_smoke_history(obj)
                                    row.append(value)
                                elif field == 'videography':
                                    value = self.filter_none(obj, field)
                                    value = self.videography_map.get(value)
                                    row.append(value)
                                elif field == 'stage':
                                    value = self.filter_none(obj, field)
                                    value = self.stage_map.get(value)
                                    row.append(value)
                                elif field == 'cStage' or field == 'pStage':
                                    stage_array = self.format_stage(obj, field)
                                    row.extend(stage_array)
                                elif (field == 'bioMet' or field == 'traSite' or
                                field == 'basDisHis' or field == 'infDisHis' or
                                field == 'tumHis' or field == 'tumFamHis' or
                                field == 'cliniManifest' or field == 'part' or
                                field == 'surSco' or field == 'lymDis' or field == 'radSite'):
                                    value = self.format_radio_data(obj, field)
                                    row.append(value)
                                elif field == 'patDia':
                                    value = self.format_patDia(obj)
                                    row.append(value)
                                elif field == 'isTre':
                                    value = self.filter_none(obj, field)
                                    value = self.isTre_map.get(value)
                                    row.append(value)
                                elif field == 'radDose':
                                    value_radDose = self.filter_none(obj, field)
                                    value_dosUnit = self.radosUnit_map.get(self.filter_none(obj, 'dosUnit'))
                                    value = str(value_radDose) + value_dosUnit
                                    row.append(value)
                                elif field == 'splTim':
                                    value_splTim = self.filter_none(obj, field)
                                    value_method = self.filter_none(obj, 'method')
                                    value = str(value_splTim) + value_method
                                    row.append(value)
                                elif field == 'isExe':
                                    value = self.filter_none(obj, field)
                                    value = self.isExe_map.get(value)
                                    row.append(value)
                                elif type(x) == bool:
                                    value = self.filter_none(self.change_bool_to_yes_or_no(x))
                                    row.append(value)
                                elif type(x) == dict:
                                    value = str(x) if x else '/'
                                    row.append(value)
                                elif type(x) == list:
                                    value = str(x) if x else '/'
                                    row.append(value)
                                else:
                                    value = self.filter_none(obj, field)
                                    row.append(value)

                        else:
                            if name == 'IniDiaPro':
                                for j in range(0, len(self.columns[i])):
                                    if j == 'pStage' or j == 'cStage':
                                        row.extend(['/','/','/'])
                                    else:
                                        row.append('/')
                            elif name == 'Lung':
                                for j in range(0, len(self.columns[i])):
                                    if j in self.lung_fields:
                                        row.extend(['/', '/', '/'])
                                    else:
                                        row.append('/')
                            else:
                                for j in range(0, len(self.columns[i])):
                                    row.append('/')

                        i += 1
                    ws.append(row)
            else:
                i = 0
                row = []
                for obj_class_name in self.tables:
                    name = obj_class_name.__name__
                    if name == 'Patient':
                        obj_array = obj_class_name.query.filter(obj_class_name.id == pid,
                                                                obj_class_name.is_delete == 0).all()
                    elif name == 'PastHis' or name == 'IniDiaPro':
                        obj_array = obj_class_name.query.filter(obj_class_name.pid == pid,
                                                                obj_class_name.is_delete == 0).all()
                    else:
                        obj_array = obj_class_name.query.filter(obj_class_name.pid == pid,
                                                                obj_class_name.treNum == treNum,
                                                                obj_class_name.is_delete == 0).all()

                    if obj_array:
                        obj = obj_array[0]

                        for field in self.columns[i]:
                            if field != 'age' and not (field in self.lung_fields):
                                x = getattr(obj, field)

                            if name == 'Immunohis':
                                if field == 'Ki67' or field == 'other' or field == 'filePath':
                                    value = self.filter_none(obj, field)
                                    row.append(value)
                                else:
                                    value = self.filter_none(obj, field)
                                    value = self.immunohis_map.get(value)
                                    row.append(value)
                            elif name == 'MoleDetec':
                                if field[-6:] == 'DetMed':
                                    value = self.filter_none(obj, field)
                                    value = self.moleDetec_DetMed_map.get(value)
                                    row.append(value)
                                elif field == 'MSI':
                                    value = self.filter_none(obj, field)
                                    value = self.moleDetec_MSI_map.get(value)
                                    row.append(value)
                                elif field in self.moleDetec_fields:
                                    value = self.filter_none(obj, field)
                                    value = self.moleDetec_map.get(value)
                                    row.append(value)
                                else:
                                    value = self.filter_none(obj, field)
                                    row.append(value)
                            elif name in self.lab_inspectation_table or name == 'Lung':
                                if field[-3:] == 'Mea':
                                    value = self.filter_none(obj, field)
                                    value = self.Mea_map.get(value)
                                    row.append(value)
                                elif field in self.lung_fields:
                                    value_exp = self.filter_none(obj, field + '_exp')
                                    value_best = self.filter_none(obj, field + '_best')
                                    value_ratio = self.filter_none(obj, field + '_ratio')
                                    row.extend([value_exp, value_best, value_ratio])
                                else:
                                    value = self.filter_none(obj, field)
                                    row.append(value)
                            elif field == 'gender':
                                value = self.filter_none(obj, field)
                                value = self.gender_map.get(value)
                                row.append(value)
                            elif field == 'age':
                                idNumber = getattr(obj, 'idNumber')
                                age = get_age_by_birth(get_birth_date_by_id_card(idNumber))
                                value = age if age else '/'
                                row.append(value)
                            elif field == 'drinkingHis':
                                value = self.format_drink_history(obj)
                                row.append(value)
                            elif field == 'smokingHis':
                                value = self.format_smoke_history(obj)
                                row.append(value)
                            elif field == 'videography':
                                value = self.filter_none(obj, field)
                                value = self.videography_map.get(value)
                                row.append(value)
                            elif field == 'stage':
                                value = self.filter_none(obj, field)
                                value = self.stage_map.get(value)
                                row.append(value)
                            elif field == 'cStage' or field == 'pStage':
                                stage_array = self.format_stage(obj, field)
                                row.extend(stage_array)
                            elif (field == 'bioMet' or field == 'traSite' or
                            field == 'basDisHis' or field == 'infDisHis' or
                            field == 'tumHis' or field == 'tumFamHis' or
                            field == 'cliniManifest' or field == 'part'):
                                value = self.format_radio_data(obj, field)
                                row.append(value)
                            elif field == 'patDia':
                                value = self.format_patDia(obj)
                                row.append(value)

                            elif type(x) == bool:
                                value = self.filter_none(self.change_bool_to_yes_or_no(x))
                                row.append(value)
                            elif type(x) == dict:
                                value = str(x) if x else '/'
                                row.append(value)
                            elif type(x) == list:
                                value = str(x) if x else '/'
                                row.append(value)
                            else:
                                value = self.filter_none(obj, field)
                                row.append(value)

                    else:
                        if name == 'IniDiaPro':
                            for j in range(0, len(self.columns[i])):
                                if j == 'pStage' or j == 'cStage':
                                    row.extend(['/', '/', '/'])
                                else:
                                    row.append('/')
                        elif name == 'Lung':
                            for j in range(0, len(self.columns[i])):
                                if j in self.lung_fields:
                                    row.extend(['/', '/', '/'])
                                else:
                                    row.append('/')
                        else:
                            for j in range(0, len(self.columns[i])):
                                row.append('/')

                    i += 1
                ws.append(row)


        wb.save('test.xlsx')
        content = save_virtual_workbook(wb)
        resp = make_response(content)
        resp.headers["Content-Disposition"] = 'attachment; filename=samples.xlsx'
        resp.headers['Content-Type'] = 'application/x-xlsx'
        return resp

    def change_bool_to_yes_or_no(self, bool_value):
        if bool_value is None:
            return '/'
        return '是' if bool_value else '否'

    def filter_none(self, data, item=None):
        if data is None:
            return '/'
        if item:
            val = getattr(data, item)
        else:
            val = data
        return val if val is not None else '/'

    def format_smoke_history(self, object):
        if object is None:
            return '/'
        is_smoke = object.smoke
        info = object.smokingHis
        value = '/'
        if is_smoke == 1:
            value = '累积吸烟时间（年）' + str(info.get('smokeYearAvg')) + '日平均吸烟量（支）' + str(info.get('smokeDayAvg'))
        elif is_smoke == 0:
            value = '无'
        return value

    def format_drink_history(self, object):
        if object is None:
            return '/'
        is_drink = object.drink
        info = object.drinkingHis
        value = '/'
        if is_drink == 1:
            value = '累积饮酒时间（年）' + str(info.get('drinkYearAvg')) + '日平均饮酒量（mL）' + str(info.get('drinkDayAvg'))
        elif is_drink == 0:
            value = '无'
        return value

    def format_radio_data(self, object, item):
        if object is None:
            return '/'
        data = getattr(object, item)
        if data is None:
            return '/'
        radio = data.get('radio')
        other = data.get('other')
        value = ",".join(radio)
        if other:
            value = value + "," + "其他: " + other
        return value

    def format_patDia(self, object):
        if object is None or object.patDia is None:
            return '/'
        patDia = object.patDia
        radio = patDia.get('radio')
        data = []
        for item in radio:
            if item != "0-5":
                data.append(self.patDia_map.get(item))
            else:
                data.append("其他：" + str(patDia.get('other')))
        return ','.join(data)

    def format_stage(self, object, item):
        value = ['/'] * 3
        if object is None:
            return value
        data = getattr(object, item)
        if data is None or data == "":
            return value
        ret = data.split(',')
        for i in range(0, 3):
            value[i] = self.filter_none(ret[i])
        return value
